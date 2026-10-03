"""Read-only XLSX extraction and whitelisted formula replay; never execute cells."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r"F:\wx\wenjian\xwechat_files\wxid_l89br9o85qih21_3b01\temp\RWTemp\2026-10\b5f79ed1eaf900d31b3438715d5a49fa\投资测算（清远酒店）.xlsx")
EXPECTED = "49b5af7ae8655bb10a505b01da17da1db995314420ad6a45de7a90cac51f94f7"
TOKEN = re.compile(r'\s*(\$?[A-Z]{1,3}\$?\d+|\d+(?:\.\d+)?|"[^"]*"|[A-Z][A-Z0-9_]*|<=|>=|<>|[+*/(),:><=-])')


class FormulaError(Exception):
    pass


class Parser:
    def __init__(self, formula):
        self.tokens = []
        pos = 0
        expression = formula[1:]
        while pos < len(expression):
            m = TOKEN.match(expression, pos)
            if not m:
                raise FormulaError(f"Unsupported expression at {expression[pos:]}")
            self.tokens.append(m.group(1))
            pos = m.end()
        self.i = 0

    def peek(self):
        return self.tokens[self.i] if self.i < len(self.tokens) else None

    def take(self, expected=None):
        t = self.peek()
        if t is None or (expected is not None and t != expected):
            raise FormulaError(f"Expected {expected}, got {t}")
        self.i += 1
        return t

    def expr(self, level=0):
        groups = [{">", "<", "<=", ">=", "=", "<>"}, {"+", "-"}, {"*", "/"}]
        if level < 3:
            node = self.expr(level + 1)
            while self.peek() in groups[level]:
                op = self.take()
                node = ("op", op, node, self.expr(level + 1))
            return node
        t = self.take()
        if t in {"+", "-"}:
            return ("unary", t, self.expr(3))
        if t == "(":
            node = self.expr()
            self.take(")")
            return node
        if t.startswith('"'):
            return ("literal", t[1:-1])
        if re.fullmatch(r"\d+(?:\.\d+)?", t):
            return ("literal", float(t))
        if re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", t):
            ref = t.replace("$", "")
            if self.peek() == ":":
                self.take(":")
                return ("range", ref, self.take().replace("$", ""))
            return ("ref", ref)
        if t not in {"SUM", "ABS", "ROUND", "IF", "IFERROR"}:
            raise FormulaError(f"Unsupported function {t}")
        self.take("(")
        args = [self.expr()]
        while self.peek() == ",":
            self.take(",")
            args.append(self.expr())
        self.take(")")
        return ("call", t, args)

    def parse(self):
        ast = self.expr()
        if self.peek() is not None:
            raise FormulaError("Trailing tokens")
        return ast


class Replay:
    def __init__(self, sheet, overrides=None):
        self.sheet = sheet
        self.overrides = overrides or {}
        self.memo = {}
        self.active = set()
        self.blank_refs = set()

    def cell(self, ref):
        if ref in self.overrides:
            return self.overrides[ref]
        if ref in self.memo:
            return self.memo[ref]
        if ref in self.active:
            raise FormulaError(f"Circular reference {ref}")
        self.active.add(ref)
        raw = self.sheet[ref].value
        if raw is None:
            self.blank_refs.add(ref)
            result = 0  # Excel blank arithmetic semantics only; not a supplied business fact.
        elif self.sheet[ref].data_type == "f":
            result = self.run(Parser(raw).parse())
        else:
            result = raw
        self.active.remove(ref)
        self.memo[ref] = result
        return result

    def run(self, node):
        kind = node[0]
        if kind == "literal":
            return node[1]
        if kind == "ref":
            return self.cell(node[1])
        if kind == "range":
            c1, r1, c2, r2 = range_boundaries(node[1] + ":" + node[2])
            return [self.cell(f"{get_column_letter(c)}{r}") for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
        if kind == "unary":
            return self.run(node[2]) * (-1 if node[1] == "-" else 1)
        if kind == "op":
            a, b = self.run(node[2]), self.run(node[3])
            op = node[1]
            if op == "+": return a + b
            if op == "-": return a - b
            if op == "*": return a * b
            if op == "/": return a / b
            if op == ">": return a > b
            if op == "<": return a < b
            if op == "<=": return a <= b
            if op == ">=": return a >= b
            if op == "=": return a == b
            if op == "<>": return a != b
        if kind == "call":
            fn, args = node[1], node[2]
            if fn == "IF":
                return self.run(args[1] if self.run(args[0]) else args[2])
            if fn == "IFERROR":
                try:
                    return self.run(args[0])
                except (FormulaError, ZeroDivisionError, TypeError, ValueError):
                    return self.run(args[1])
            vals = [self.run(a) for a in args]
            if fn == "SUM":
                flat = [x for v in vals for x in (v if isinstance(v, list) else [v])]
                return sum(x for x in flat if isinstance(x, (int, float)))
            if fn == "ABS": return abs(vals[0])
            if fn == "ROUND":
                x, digits = vals
                scale = 10 ** int(digits)
                return math.copysign(math.floor(abs(x) * scale + 0.5) / scale, x)
        raise FormulaError(f"Unsupported node {node}")


def save_json(name, obj):
    (ROOT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def payback(initial, flows):
    balance = -initial
    for year, flow in enumerate(flows, 1):
        before = balance
        balance += flow
        if before < 0 <= balance and flow > 0:
            return year - 1 + (-before / flow)
    return None


def main():
    source = SOURCE if SOURCE.exists() else ROOT / "source.xlsx"
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != EXPECTED:
        raise ValueError("Source fingerprint changed")
    if source != ROOT / "source.xlsx":
        dest = ROOT / "source.xlsx"
        if dest.exists() and hashlib.sha256(dest.read_bytes()).hexdigest() != EXPECTED:
            raise ValueError("Preserved source collision")
        if not dest.exists():
            shutil.copy2(source, dest)
    book = openpyxl.load_workbook(source, data_only=False)
    caches = openpyxl.load_workbook(source, data_only=True)
    sheet = book.active
    replay = Replay(sheet)
    cells, formulas = [], []
    for row in sheet:
        for cell in row:
            if cell.value is None:
                continue
            cached = caches[sheet.title][cell.coordinate].value
            record = {"sheet": sheet.title, "cell": cell.coordinate, "raw": cell.value, "cached": cached, "number_format": cell.number_format, "hidden_row": bool(sheet.row_dimensions[cell.row].hidden)}
            cells.append(record)
            if cell.data_type == "f":
                actual = replay.cell(cell.coordinate)
                match = math.isclose(actual, cached, rel_tol=1e-12, abs_tol=1e-6) if isinstance(actual, (int, float)) and isinstance(cached, (int, float)) else actual == cached
                formulas.append({"cell": cell.coordinate, "formula": cell.value, "cached": cached, "replayed": actual, "matched": match})
    save_json("cells.json", cells)
    save_json("formula-replay.json", formulas)
    yearly = []
    cashflows, profits = [], []
    for col in range(3, 13):
        c = get_column_letter(col)
        r = {"year": replay.cell(f"{c}16"), "revpar_yuan": replay.cell(f"{c}17"), "room_revenue_yuan": replay.cell(f"{c}18"), "operating_cost_yuan": replay.cell(f"{c}19"), "gop_yuan": replay.cell(f"{c}20"), "rent_yuan": replay.cell(f"{c}21"), "depreciation_yuan": replay.cell(f"{c}23"), "management_fee_yuan": replay.cell(f"{c}24"), "source_pretax_profit_yuan": replay.cell(f"{c}25"), "source_cumulative_profit_proxy_yuan": replay.cell(f"{c}26")}
        r["pretax_cash_proxy_yuan"] = r["source_pretax_profit_yuan"] + r["depreciation_yuan"]
        profits.append(r["source_pretax_profit_yuan"])
        cashflows.append(r["pretax_cash_proxy_yuan"])
        yearly.append(r)
    initial = replay.cell("L13") * 10000
    with (ROOT / "annual-replay.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(yearly[0]))
        writer.writeheader()
        writer.writerows(yearly)
    scenarios = []
    for adr in [180, 200, 220, 250, 280, 300]:
        for occ in [0.50, 0.60, 0.70, 0.80, 0.85, 0.90]:
            rr = Replay(sheet, {"G2": adr, "G3": occ, "G4": occ})
            pf = [rr.cell(f"{get_column_letter(c)}25") for c in range(3, 13)]
            cf = [pf[i] + rr.cell(f"{get_column_letter(i+3)}23") for i in range(10)]
            scenarios.append({"adr_yuan": adr, "occupancy": occ, "first_year_pretax_profit_yuan": pf[0], "first_year_pretax_cash_proxy_yuan": cf[0], "profit_proxy_payback_years": payback(initial, pf), "pretax_cash_proxy_payback_years": payback(initial, cf), "boundary": "synthetic_sensitivity; occupancy_same_all_years; source_year3_growth; excludes_tax_financing_capex_working_capital"})
    with (ROOT / "sensitivity.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(scenarios[0]))
        writer.writeheader()
        writer.writerows(scenarios)
    term_override = Replay(sheet, {"B5": 5})
    rent_override = Replay(sheet, {"C22": 0.1, "D22": 0.1})
    construction_override = Replay(sheet, {"G13": 6})
    no_cross = Replay(sheet, {"G2": 180, "G3": .5, "G4": .5})
    no_cross_flows = [no_cross.cell(f"{get_column_letter(c)}25") for c in range(3, 13)]
    zero_return = replay.run(Parser('=IFERROR(IF(-3371300/0>0,"N/A",0+ABS(-3371300)/3371300),"N/A")').parse())
    manifest = {
        "source_name": "投资测算（清远酒店）.xlsx", "source_original_path": str(SOURCE), "source_sha256": digest,
        "source_preserved": "source.xlsx", "source_created_metadata": str(book.properties.created), "source_modified_metadata": str(book.properties.modified),
        "processing_date_asia_shanghai": "2026-10-01", "source_business_date": None, "hotel_id": None, "tenant_id": None, "platform": None,
        "quality_status": "unverified_source_assumptions", "reference_only": True, "decision_safe": False, "task_draft_safe": False, "external_write_authorized": False,
        "task_mode": "classify", "source_instructions_active": False, "source_version": "sha256:" + digest, "sheet": sheet.title,
        "nonempty_cells": len(cells), "formula_count": len(formulas), "hidden_rows": [i for i, r in sheet.row_dimensions.items() if r.hidden],
        "sheets": book.sheetnames, "external_links": len(book._external_links), "macros": False,
        "interpretation_boundary": "Workbook scenario assumptions, no verified hotel/business/account facts. Blank references replay as Excel arithmetic zero only and remain missing in knowledge.",
    }
    save_json("source-manifest.json", manifest)
    result = {
        "formula_count": len(formulas), "cache_matches": sum(x["matched"] for x in formulas), "mismatches": [x for x in formulas if not x["matched"]],
        "blank_references_used_excel_semantics_only": sorted(replay.blank_refs), "initial_investment_yuan": initial,
        "source_payback_years": replay.cell("M27"), "unrounded_profit_proxy_payback_years": payback(initial, profits),
        "pretax_cash_proxy_payback_years": payback(initial, cashflows), "source_total_profit_yuan": replay.cell("M25"),
        "sum_annual_profit_yuan": sum(profits), "source_total_overstatement_yuan": replay.cell("M25") - sum(profits),
        "source_total_depreciation_yuan": replay.cell("M23"), "sum_annual_depreciation_yuan": sum(x["depreciation_yuan"] for x in yearly),
        "source_seventh_year_rent_escalation_input": replay.cell("I22"),
        "break_even_first_year_profit_occupancy": (replay.cell("C19") + replay.cell("C21") + replay.cell("C23")) / (153 * 365 * 250 * .965),
        "break_even_first_year_cash_proxy_occupancy": (replay.cell("C19") + replay.cell("C21")) / (153 * 365 * 250 * .965),
        "break_even_first_year_profit_adr_at_70pct": (replay.cell("C19") + replay.cell("C21") + replay.cell("C23")) / (153 * 365 * .7 * .965),
        "decorating_unit_cost_implied_yuan": replay.cell("G12") * replay.cell("G11"), "manual_decorating_yuan": replay.cell("L11") * 10000,
        "first_year_occupied_night_cost_if_alternate_yuan": 100 * 153 * 365 * .7,
        "first_year_occupied_vs_available_cost_delta_yuan": replay.cell("C19") - 100 * 153 * 365 * .7,
        "source_no_payback_scenario": {"adr_yuan": 180, "occupancy_all_years": .5, "source_M27_years": no_cross.cell("M27"), "first_crossing_years": payback(initial, no_cross_flows)},
        "checks": {
            "all_formula_caches_replayed": all(x["matched"] for x in formulas),
            "five_year_term_input_does_not_shorten_horizon": term_override.cell("L16") == 10 and term_override.cell("M25") == replay.cell("M25"),
            "rent_escalation_row_not_applied": rent_override.cell("D21") == replay.cell("D21") and rent_override.cell("D25") == replay.cell("D25"),
            "construction_six_month_rent_computed_but_not_cashflow": construction_override.cell("B21") == 1377000 and construction_override.cell("B26") == replay.cell("B26"),
            "exact_zero_crossing_source_formula_returns_na": zero_return == "N/A",
            "source_no_crossing_M27_returns_zero_misleadingly": no_cross.cell("M27") == 0 and payback(initial, no_cross_flows) is None,
            "helper_no_crossing_scenario_returns_none": payback(initial, [-100] * 10) is None,
        },
        "scenario_boundary": "Sensitivity and cash proxies are arithmetic demonstrations, not investment advice, actual market forecasts, verified cashflows or source author confirmed corrections.",
    }
    save_json("analysis-results.json", result)
    print(json.dumps({"manifest": manifest, "results": result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
